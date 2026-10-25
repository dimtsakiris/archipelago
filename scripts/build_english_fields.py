#!/usr/bin/env python3
"""
build_english_fields.py — generate QuickStatements that add ENGLISH values to the
parallel properties (Creator EN, Place EN, Keywords EN, Notes EN, AN project EN).

Mining items: English values come from the English source workbook.
Amorgos items: no English source exists, so values come from the small dictionaries
below (places, creator names = transliteration, keywords). REVIEW them.

    python build_english_fields.py --qid-map qid_map.csv \
        --el digital_cyclades_extraction_greek.xlsx --en digital_cyclades_extraction_english.xlsx \
        --amorgos-csv am/items_clean.csv --out-dir out
"""
import argparse, csv, re, sys
from pathlib import Path
import openpyxl

# >>> property IDs of the new parallel properties (change if Wikibase gave other numbers)
PID = {"creator": "P17", "place": "P18", "keywords": "P19", "notes": "P20", "an_project": "P21"}

AMORGOS_COLLECTIONS = {"cmt_amorgos_archival_collection_hozoviotissa": "Q3",
                       "cmt_amorgos_archival_collection_ioannis i vekris": "Q4",
                       "cmt_amorgos_archival_collection_skopelitis": "Q5"}

PLACES = {"Αμοργός": "Amorgos", "Αιγιάλη": "Aigiali", "Ίος": "Ios", "Νικουριά": "Nikouria",
          "Μικρές Κυκλάδες": "Small Cyclades", "Νάξος": "Naxos", "Κουφονήσια": "Koufonisia",
          "Κουφονήσι": "Koufonisi", "Αντικέρι": "Antikeri", "Λεβρεσό": "Levreso", "Κέρος": "Keros",
          "Σύρος": "Syros", "Κυκλάδες": "Cyclades", "Σπέτσες": "Spetses", "Ηρακλειά": "Iraklia"}
CREATORS = {"Γιάννης Σκοπελίτης": "Giannis Skopelitis", "Μιχάλης Βεκρής": "Michalis Bekris",
            "Μιχάλης Κρόσμαν": "Michalis Krosman", "Δημήτρης Βεκρής": "Dimitris Bekris",
            "Δημήτρης Γαβαλάς": "Dimitris Gavalas", "Ιωάννης Βεκρής": "Ioannis Bekris",
            "Ιωάννης Η. Βεκρής": "Ioannis H. Bekris", "Νικόλαος Βεκρής": "Nikolaos Bekris",
            "Παρασκευάς Ρούσσος": "Paraskevas Roussos", "Ιωάννης Ψακής": "Ioannis Psakis",
            "Γεώργιος Ρούσσος": "Georgios Roussos", "Παρασκευώ Θεολογίτου": "Paraskevo Theologitou",
            "Επαγγελματικός Αλιευτικός Σύλλογος Αμοργού": "Professional Fishermen's Association of Amorgos"}
KEYWORDS = {"καΐκι": "caique", "καϊκι": "caique", "ψαράς": "fisherman", "θάλασσα": "sea",
 "Κατάπολα": "Katapola", "Αιγιάλη": "Aigiali", "ψάρια": "fish", "πλοίο": "ship", "δίχτυα": "nets",
 "Χώρα": "Chora", "αλιεία": "fishing", "Δημήτρης Βεκρής": "Dimitris Bekris",
 "Κώστας Σιγάλας": "Kostas Sigalas", "ψάρεμα": "fishing", "Γιάννης Βεκρής": "Giannis Bekris",
 "γρι-γρι": "purse seine (gri-gri)", "Ίος": "Ios", "Πανορμίτης": "Panormitis", "ψαράδες": "fishermen",
 "Σκοπελίτης": "Skopelitis", "αλίευμα": "catch", "Αιγιαλή": "Aigiali", "Μιχάλης Σιμιδαλάς": "Michalis Simidalas",
 "τράτα": "beach seine (trata)", "καθάρισμα δίχτυα": "net cleaning", "Καμάρι": "Kamari", "Μαριάννα": "Marianna",
 "αλιείς": "fishermen", "Λεβρεσό": "Levreso", "καθάρισμα διχτυών": "net cleaning", "ρίξιμο διχτυών": "casting nets",
 "Νικουριά": "Nikouria", "Γιώργης ο κουρέας": "Giorgis the barber", "Νικολότασσος": "Nikolotassos",
 "χταπόδια": "octopuses", "καμάκι": "fishing spear", "Νικόλαος Ψακής": "Nikolaos Psakis",
 "Ουρανία Ψακή": "Ourania Psaki", "Δημήτρης Γαβαλάς": "Dimitris Gavalas", "Σ´υρο": "Syros",
 "Μιχάλης Βεκρής": "Michalis Bekris", "Βεκρής": "Bekris", "ΑΙγιάλη": "Aigiali", "Πανόραμα": "panorama",
 "ψαραδες": "fishermen", "οικογένεια": "family", "ηλιοβασίλεμα": "sunset", "μπουρού": "conch horn (bourou)",
 "Αγία Άννα": "Agia Anna", "διαλογή ψαριών": "fish sorting", "Καλοταρίτισσα": "Kalotaritissa", "κάρτα": "card",
 "ευχές": "greetings", "Άγιος Γεώργιος Βαλσαμίτης": "Agios Georgios Valsamitis", "Νάξος": "Naxos",
 "Δημήτρης Σκοπελίτης": "Dimitris Skopelitis", "γλέντι": "celebration", "βιολί": "violin", "Κουφονήσι": "Koufonisi",
 "Χοζοβιώτισσα": "Hozoviotissa", "Σποράδες": "Sporades", "Μικρές Κυκλάδες": "Small Cyclades",
 "Δημήτριος": "Dimitrios", "Σπέτσες": "Spetses", "Μιχάλης Ψακής": "Michalis Psakis", "Ηρακλεία": "Iraklia",
 "εμπορεύματα": "goods"}
AN_EN = "Extraction in the Cyclades: Digital Traces"


def q(v):
    return '"' + str(v).replace('"', "'").replace("\n", " ").strip() + '"'


def norm_kw(k):
    return re.sub(r"\s+", " ", k.replace('"', "").replace("'", "").strip())


def note_en(raw):
    s = re.sub(r"\s+", " ", str(raw).strip())
    return re.sub(r"^(περ\.?|γ\.?|g\.?|c\.?|circa)\s*", "c. ", s, flags=re.I)


def has_note(raw):
    """Same rule as build_import.parse_date: a P11 note exists unless the date is a plain year."""
    if raw is None or str(raw).strip() == "":
        return False
    try:
        int(float(str(raw).strip())); return False
    except ValueError:
        return bool(re.search(r"\d{4}", str(raw)))


def rows(ws):
    for r in range(3, ws.max_row + 1):
        row = list(ws.iter_rows(min_row=r, max_row=r, values_only=True))[0]
        if row and row[1] is not None:
            yield row


def c(v):
    return "" if v is None else str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--qid-map", required=True); ap.add_argument("--el", required=True)
    ap.add_argument("--en", required=True); ap.add_argument("--amorgos-csv", required=True)
    ap.add_argument("--out-dir", default="out"); ap.add_argument("--batch-size", type=int, default=300)
    a = ap.parse_args()
    qmap = {r["id"]: r["item"].rsplit("/", 1)[-1] for r in csv.DictReader(open(a.qid_map, encoding="utf-8"))}
    lines, review, missing = [], [], []

    def add(qid, key, val):
        if val:
            lines.append(f"{qid}|{PID[key]}|{q(val)}")

    # ---- mining items + collections (English workbook)
    wel, wen = openpyxl.load_workbook(a.el, data_only=True), openpyxl.load_workbook(a.en, data_only=True)
    for el, en in zip(rows(wel["ITEM"]), rows(wen["ITEM"])):
        ident = c(el[2]); qid = qmap.get(ident)
        if not qid: missing.append(ident); continue
        add(qid, "creator", c(en[4])); add(qid, "place", c(en[7])); add(qid, "keywords", c(en[14]))
        add(qid, "an_project", c(en[15]))
        if has_note(el[6]): add(qid, "notes", note_en(c(en[6]) or el[6]))
    for el, en in zip(rows(wel["COLLECTION"]), rows(wen["COLLECTION"])):
        qid = qmap.get(c(el[2]))
        if not qid: missing.append(c(el[2])); continue
        add(qid, "creator", c(en[4])); add(qid, "keywords", c(en[6])); add(qid, "an_project", c(en[7]))

    # ---- Amorgos items + collections (dictionaries)
    for r in csv.DictReader(open(a.amorgos_csv, encoding="utf-8")):
        qid = qmap.get(r["identifier"])
        if not qid: missing.append(r["identifier"]); continue
        cr, pl = r["creator"].strip(), r["place"].strip()
        if cr:
            add(qid, "creator", CREATORS.get(cr)); 
            if cr not in CREATORS: review.append(("creator", cr))
        if pl:
            add(qid, "place", PLACES.get(pl))
            if pl not in PLACES: review.append(("place", pl))
        kws = []
        for k in r["keywords"].split(","):
            k0 = norm_kw(k)
            if not k0: continue
            if k0 in KEYWORDS: kws.append(KEYWORDS[k0])
            else: review.append(("keyword", k0))
        if kws: add(qid, "keywords", ", ".join(dict.fromkeys(kws)))
        if has_note(r["date"]): add(qid, "notes", note_en(r["date"]))
    for ident, qid in AMORGOS_COLLECTIONS.items():
        cr = {"cmt_amorgos_archival_collection_hozoviotissa": "Επαγγελματικός Αλιευτικός Σύλλογος Αμοργού",
              "cmt_amorgos_archival_collection_ioannis i vekris": "Ιωάννης Η. Βεκρής",
              "cmt_amorgos_archival_collection_skopelitis": "Γιάννης Σκοπελίτης"}[ident]
        add(qid, "creator", CREATORS[cr])

    out = Path(a.out_dir); out.mkdir(parents=True, exist_ok=True)
    n = a.batch_size
    for i in range(0, len(lines), n):
        (out / f"batch_en_{i // n + 1:02d}.qs").write_text("\n".join(lines[i:i + n]) + "\n", encoding="utf-8")
    (out / "review_untranslated.txt").write_text("\n".join(f"{k}: {v}" for k, v in dict.fromkeys(review)), encoding="utf-8")
    print(f"{len(lines)} statements -> {-(-len(lines)//n)} batches; missing Q-ID: {len(missing)}; to review: {len(set(review))}")
    print("missing:", missing[:5])


if __name__ == "__main__":
    main()
